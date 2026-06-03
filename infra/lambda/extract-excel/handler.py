"""Excel データ抽出 Lambda: openpyxl でセル値を読み取り + XML 解析でテキストボックスを抽出"""

import json
import os
import io
import zipfile
import re
import boto3

s3 = boto3.client("s3")


def _extract_textboxes(xlsx_bytes: bytes) -> list[dict]:
    """xlsx の XML からテキストボックスの内容を抽出する"""
    textboxes = []
    with zipfile.ZipFile(io.BytesIO(xlsx_bytes)) as zf:
        for name in zf.namelist():
            if not name.startswith("xl/drawings/"):
                continue
            xml_content = zf.read(name).decode("utf-8", errors="replace")
            texts = re.findall(r"<a:t>([^<]+)</a:t>", xml_content)
            if texts:
                textboxes.append({"source": name, "texts": texts})
    return textboxes


def handler(event, context):
    bucket = event["bucket"]
    key = event["key"]
    report_id = event["report_id"]

    obj = s3.get_object(Bucket=bucket, Key=key)
    body = obj["Body"].read()

    try:
        import openpyxl

        wb = openpyxl.load_workbook(io.BytesIO(body), data_only=True)
    except ImportError:
        return {
            "bucket": bucket,
            "key": key,
            "report_id": report_id,
            "error": "openpyxl not available",
            "extractionType": "excel",
        }

    sheets_data = []
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        rows = []
        for row in ws.iter_rows(values_only=False):
            row_data = []
            for cell in row:
                row_data.append(
                    {
                        "coordinate": cell.coordinate,
                        "value": str(cell.value) if cell.value is not None else None,
                    }
                )
            rows.append(row_data)
        sheets_data.append({"sheetName": sheet_name, "rows": rows})

    textboxes = _extract_textboxes(body)

    return {
        "bucket": bucket,
        "key": key,
        "report_id": report_id,
        "extractionType": "excel",
        "sheets": sheets_data,
        "textboxes": textboxes,
    }
